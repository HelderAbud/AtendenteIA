import json
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

TEMAS_NEGOCIO = ("prazo", "sla", "renovacao", "boleto", "chamado")
TEMA_DESCONHECIDO = "desconhecido"
TEMAS = TEMAS_NEGOCIO + (TEMA_DESCONHECIDO,)


class BaseConhecimento:
    def __init__(self, caminho: Path) -> None:
        if not caminho.is_file():
            raise FileNotFoundError(caminho)
        dados = json.loads(caminho.read_text(encoding="utf-8"))
        self._itens = _validar(dados)

    def obter(self, tema: str) -> dict:
        return self._itens[tema]


def _validar(dados: object) -> dict[str, dict]:
    if not isinstance(dados, dict):
        raise ValueError("base invalida")
    itens: dict[str, dict] = {}
    for tema in TEMAS:
        if tema not in dados or not isinstance(dados[tema], dict):
            raise ValueError(f"tema ausente: {tema}")
        resposta = dados[tema].get("resposta")
        encaminhar = dados[tema].get("encaminhar_humano")
        if not isinstance(resposta, str) or not resposta.strip():
            raise ValueError(f"resposta ausente: {tema}")
        if not isinstance(encaminhar, bool):
            raise ValueError(f"encaminhar ausente: {tema}")
        if encaminhar is not (tema == TEMA_DESCONHECIDO):
            raise ValueError(f"encaminhar inesperado: {tema}")
        itens[tema] = {
            "resposta": resposta,
            "encaminhar_humano": encaminhar,
        }
    return itens


class BasePdf:
    def __init__(self, caminho: Path) -> None:
        if not caminho.is_file():
            raise FileNotFoundError(caminho)
        self._itens = _validar(_itens_do_pdf(caminho))

    def obter(self, tema: str) -> dict:
        return self._itens[tema]


def _itens_do_pdf(caminho: Path) -> dict[str, dict]:
    leitor = PdfReader(caminho)
    texto = "\n".join((pagina.extract_text() or "") for pagina in leitor.pages)
    blocos: dict[str, list[str]] = {}
    atual: str | None = None
    for linha in texto.splitlines():
        nome = linha.strip().lower()
        if nome in TEMAS:
            if nome in blocos:
                raise ValueError(f"tema repetido: {nome}")
            blocos[nome] = []
            atual = nome
            continue
        if atual is not None and linha.strip():
            blocos[atual].append(linha.strip())
    itens: dict[str, dict] = {}
    for tema in TEMAS:
        if tema not in blocos:
            raise ValueError(f"tema ausente: {tema}")
        resposta = " ".join(blocos[tema]).strip()
        if not resposta:
            raise ValueError(f"resposta ausente: {tema}")
        itens[tema] = {
            "resposta": resposta,
            "encaminhar_humano": tema == TEMA_DESCONHECIDO,
        }
    return itens


def _celula(valor: object) -> str:
    if valor is None:
        return ""
    return str(valor).strip()


class BasePlanilha:
    def __init__(self, caminho: Path) -> None:
        if not caminho.is_file():
            raise FileNotFoundError(caminho)
        self._itens = _validar(_itens_da_planilha(caminho))

    def obter(self, tema: str) -> dict:
        return self._itens[tema]


def _itens_da_planilha(caminho: Path) -> dict[str, dict]:
    livro = load_workbook(caminho, data_only=False)
    aba = livro.worksheets[0]
    linhas = list(aba.iter_rows(values_only=True))
    if not linhas:
        raise ValueError("cabecalho ausente")
    cabecalho = [_celula(valor).lower() for valor in linhas[0]]
    if len(cabecalho) > 2 and any(cabecalho[2:]):
        raise ValueError("coluna a mais")
    if cabecalho[:2] != ["tema", "resposta"]:
        raise ValueError("cabecalho invalido")
    blocos: dict[str, str] = {}
    for linha in linhas[1:]:
        celulas = [_celula(valor) for valor in linha]
        if len(celulas) > 2 and any(celulas[2:]):
            raise ValueError("coluna a mais")
        tema = celulas[0].lower() if celulas else ""
        resposta = celulas[1] if len(celulas) > 1 else ""
        if not tema and not resposta:
            continue
        if tema not in TEMAS:
            raise ValueError(f"tema ausente: {tema}")
        if tema in blocos:
            raise ValueError(f"tema repetido: {tema}")
        if not resposta:
            raise ValueError(f"resposta ausente: {tema}")
        blocos[tema] = resposta
    itens: dict[str, dict] = {}
    for tema in TEMAS:
        if tema not in blocos:
            raise ValueError(f"tema ausente: {tema}")
        itens[tema] = {
            "resposta": blocos[tema],
            "encaminhar_humano": tema == TEMA_DESCONHECIDO,
        }
    return itens
